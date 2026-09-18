import csv
import os
import re
from docx import Document

# Absolute paths
DOC_PATH = r'C:\xampp\htdocs\exam_system\questions.docx'
CSV_PATH = r'C:\xampp\htdocs\exam_system\questions.csv'

if not os.path.exists(DOC_PATH):
    print(f"Error: {DOC_PATH} not found!")
    exit()

doc = Document(DOC_PATH)
questions_data = []

# Method 1: Check for Table Layout
if len(doc.tables) > 0:
    print("Detected Table format in document...")
    for table in doc.tables:
        for row in table.rows[1:]:  # Skip Header
            cols = [cell.text.strip() for cell in row.cells]
            if len(cols) >= 6:
                questions_data.append([
                    cols[0], cols[1], cols[2], cols[3], cols[4], 
                    cols[5].strip().upper()[0] if cols[5] else 'A'
                ])

# Method 2: Parse Plain Unstructured Text
if len(questions_data) == 0:
    print("Detected Plain Text format in document...")
    full_text = "\n".join([p.text.strip() for p in doc.paragraphs if p.text.strip()])
    
    # Split text by numbered questions (e.g., "1.", "2.")
    raw_blocks = re.split(r'\n(?=\d+[\.\)]\s*)', full_text)
    
    for block in raw_blocks:
        lines = [line.strip() for line in block.split('\n') if line.strip()]
        if not lines:
            continue
        
        q_text = lines[0]
        q_text = re.sub(r'^\d+[\.\)]\s*', '', q_text)
        
        opt_a = opt_b = opt_c = opt_d = ""
        correct_key = "A"
        
        for line in lines[1:]:
            if re.match(r'^[A][\.\)]\s*', line, re.IGNORECASE):
                opt_a = re.sub(r'^[A][\.\)]\s*', '', line, flags=re.IGNORECASE)
            elif re.match(r'^[B][\.\)]\s*', line, re.IGNORECASE):
                opt_b = re.sub(r'^[B][\.\)]\s*', '', line, flags=re.IGNORECASE)
            elif re.match(r'^[C][\.\)]\s*', line, re.IGNORECASE):
                opt_c = re.sub(r'^[C][\.\)]\s*', '', line, flags=re.IGNORECASE)
            elif re.match(r'^[D][\.\)]\s*', line, re.IGNORECASE):
                opt_d = re.sub(r'^[D][\.\)]\s*', '', line, flags=re.IGNORECASE)
        
        if q_text and opt_a and opt_b and opt_c and opt_d:
            questions_data.append([q_text, opt_a, opt_b, opt_c, opt_d, correct_key])

# Save output to CSV
with open(CSV_PATH, 'w', newline='', encoding='utf-8') as f:
    writer = csv.writer(f)
    writer.writerow(['question_text', 'option_a', 'option_b', 'option_c', 'option_d', 'correct_option'])
    writer.writerows(questions_data)

print(f"Success! Converted {len(questions_data)} questions to {CSV_PATH}")